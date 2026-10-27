from __future__ import annotations

import logging
import tempfile
import uuid
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Protocol, runtime_checkable

from app.documents.parsed import ParsedDocument, ParsedPage

logger = logging.getLogger(__name__)


@runtime_checkable
class DocumentParser(Protocol):
    def parse(self, *, path: Path, mime_type: str, document_id: uuid.UUID, title: str) -> ParsedDocument: ...


class TxtDocumentParser:
    def parse(self, *, path: Path, mime_type: str, document_id: uuid.UUID, title: str) -> ParsedDocument:
        text = path.read_text(encoding="utf-8", errors="replace")
        return ParsedDocument(
            document_id=document_id,
            title=title,
            text=text,
            pages=[ParsedPage(page_number=1, text=text)],
            metadata={"parser": "txt"},
        )


class DoclingDocumentParser:
    """Docling-backed parser with safe fallbacks for dev/test."""

    def __init__(self, *, pdf_use_pypdf: bool = False) -> None:
        self._txt = TxtDocumentParser()
        self._pdf_use_pypdf = pdf_use_pypdf

    def parse(self, *, path: Path, mime_type: str, document_id: uuid.UUID, title: str) -> ParsedDocument:
        suffix = path.suffix.lower()
        if mime_type in {"text/plain", "text/markdown"} or suffix in {".txt", ".md", ".markdown"}:
            return self._txt.parse(path=path, mime_type=mime_type, document_id=document_id, title=title)
        if self._pdf_use_pypdf and (suffix == ".pdf" or mime_type == "application/pdf"):
            return self._fallback(path=path, mime_type=mime_type, document_id=document_id, title=title)
        try:
            from docling.document_converter import DocumentConverter

            converter = DocumentConverter()
            result = converter.convert(str(path))
            doc = result.document
            text = doc.export_to_markdown() if hasattr(doc, "export_to_markdown") else str(doc)
            pages: list[ParsedPage] = []
            if hasattr(doc, "pages"):
                for i, page in enumerate(doc.pages, start=1):
                    page_text = getattr(page, "text", "") or ""
                    pages.append(ParsedPage(page_number=i, text=page_text))
            if not pages:
                pages = [ParsedPage(page_number=1, text=text)]
            return ParsedDocument(
                document_id=document_id,
                title=title,
                text=text,
                pages=pages,
                metadata={"parser": "docling"},
            )
        except Exception as exc:
            logger.warning("docling parse failed (%s) — trying fallback", type(exc).__name__)
            return self._fallback(path=path, mime_type=mime_type, document_id=document_id, title=title)

    def _fallback(self, *, path: Path, mime_type: str, document_id: uuid.UUID, title: str) -> ParsedDocument:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            pages = []
            parts: list[str] = []
            for i, page in enumerate(reader.pages, start=1):
                t = page.extract_text() or ""
                pages.append(ParsedPage(page_number=i, text=t))
                parts.append(t)
            text = "\n\n".join(parts)
            return ParsedDocument(document_id=document_id, title=title, text=text, pages=pages, metadata={"parser": "pypdf_fallback"})
        if suffix == ".docx":
            text = self._extract_docx_text(path)
            return ParsedDocument(
                document_id=document_id,
                title=title,
                text=text,
                pages=[ParsedPage(page_number=1, text=text)],
                metadata={"parser": "docx_ooxml_fallback" if text else "python_docx_fallback"},
            )
        if suffix in {".xlsx", ".xls"}:
            import openpyxl

            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            rows: list[str] = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    rows.append("\t".join("" if c is None else str(c) for c in row))
            text = "\n".join(rows)
            return ParsedDocument(
                document_id=document_id,
                title=title,
                text=text,
                pages=[ParsedPage(page_number=1, text=text)],
                metadata={"parser": "openpyxl_fallback"},
            )
        return self._txt.parse(path=path, mime_type=mime_type, document_id=document_id, title=title)

    @staticmethod
    def _extract_docx_text(path: Path) -> str:
        try:
            import docx

            document = docx.Document(str(path))
            return "\n".join(p.text for p in document.paragraphs)
        except Exception:
            pass
        try:
            with zipfile.ZipFile(path) as zf:
                raw = zf.read("word/document.xml")
        except Exception:
            return ""
        root = ET.fromstring(raw)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        parts = [node.text for node in root.findall(".//w:t", ns) if node.text]
        return "\n".join(parts)


def materialize_bytes_for_parse(data: bytes, suffix: str) -> Path:
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    tmp.write(data)
    tmp.close()
    return Path(tmp.name)


def build_document_parser(settings) -> DocumentParser:
    return DoclingDocumentParser(pdf_use_pypdf=bool(getattr(settings, "document_pdf_use_pypdf", False)))
